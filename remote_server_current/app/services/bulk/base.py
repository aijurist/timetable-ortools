"""
app/services/bulk/base.py
===========================
Shared engine for CSV / XLSX bulk import.
Provides file parsing, column mapping, type coercion, and department resolution.
"""
from __future__ import annotations

import csv
import os
import re
import tempfile
import uuid
from pathlib import Path
from typing import Any

from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.upload import ALLOWED_EXTENSIONS, MAX_BULK_ROWS, MAX_UPLOAD_SIZE
from app.models.institution import Department


# ---------------------------------------------------------------------------
# Column header aliases — case-insensitive fuzzy matched
# ---------------------------------------------------------------------------

HEADER_ALIASES: dict[str, list[str]] = {
    "name": [
        "name", "full name", "faculty name", "staff name",
        "course name", "room name", "student name", "fullname",
    ],
    "email": [
        "email", "email address", "email id", "mail", "e-mail",
    ],
    "code": [
        "code", "course code", "subject code", "room code",
        "faculty code", "department code", "subject", "course id",
        "course_code", "subject id",
    ],
    "department_id": [
        "department", "dept", "department code", "dept code",
        "department name", "dept name", "department_id",
    ],
    "capacity": [
        "capacity", "seats", "max seats", "room capacity",
    ],
    "enrollment_number": [
        "enrollment no", "enrollment number", "enrollment_no",
        "roll no", "roll number", "roll_no",
    ],
    "room_type": [
        "room type", "type", "room_type",
    ],
    "employment_type": [
        "employment type", "employment_type", "employment",
    ],
    "designation": [
        "designation", "title",
    ],
    "lecture_hours": [
        "lecture hours", "lecture", "L", "lec",
    ],
    "tutorial_hours": [
        "tutorial hours", "tutorial", "T", "tut",
    ],
    "practical_hours": [
        "practical hours", "practical", "lab hours", "P", "lab",
    ],
    "credits": [
        "credits", "credit",
    ],
    "degree_type": [
        "degree", "degree type", "degree_type",
    ],
    "program": [
        "program", "program name", "course",
    ],
    "semester": [
        "semester", "sem",
    ],
    "year_of_study": [
        "year of study", "year", "year_of_study",
    ],
    "building": [
        "building", "building name",
    ],
    "campus": [
        "campus", "campus name",
    ],
    "tags": [
        "tags", "room tags", "facilities",
    ],
    "room_tags": [
        "room tags", "room_tags",
    ],
    "availability_blacklist": [
        "availability blacklist", "availability_blacklist", "blacklist",
    ],
    "age": [
        "age",
    ],
    "elective_type": [
        "elective type", "elective", "elective_type",
        "professional elective", "pe", "open elective", "oe",
    ],
    "elective_semester": [
        "elective semester", "pe semester", "elective sem",
        "elective_semester",
    ],
    "phone": [
        "phone", "phone number", "mobile", "mobile number",
        "contact", "contact number",
    ],
    "gender": [
        "gender",
    ],
    "staff_code": [
        "staff code", "staff id", "employee code", "emp code",
        "staff_code",
    ],
    "employee_id": [
        "employee id", "emp id", "employee number", "emp number",
        "employee_id",
    ],
    "faculty_name": [
        "faculty name", "teacher name", "staff name", "teacher", "faculty",
        "faculty_name", "teacher_name",
    ],
    "section_count": [
        "section count", "sections", "section", "no of sections",
        "section_count", "num sections", "number of sections",
    ],
    "room_hint": [
        "room hint", "room", "capacity needed", "student count",
        "min capacity", "students", "room requirement", "room_hint",
        "min students", "room capacity needed",
    ],
    "requested_dept": [
        "requested dept", "request dept", "service dept", "from dept",
        "cross dept", "requested_dept", "service department",
        "cross department", "request department",
    ],
    "notes": [
        "notes", "note", "comments", "comment", "remarks", "remark",
        "additional info", "additional notes", "description",
    ],
}


# ---------------------------------------------------------------------------
# Levenshtein distance for fuzzy header matching
# ---------------------------------------------------------------------------

def _levenshtein(a: str, b: str) -> int:
    if len(a) < len(b):
        a, b = b, a
    if not b:
        return len(a)
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a):
        curr = [i + 1]
        for j, cb in enumerate(b):
            cost = 0 if ca == cb else 1
            curr.append(min(curr[j] + 1, prev[j + 1] + 1, prev[j] + cost))
        prev = curr
    return prev[-1]


# ---------------------------------------------------------------------------
# File parsing
# ---------------------------------------------------------------------------

async def save_upload_to_temp(file: UploadFile) -> str:
    ext = Path(file.filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Allowed: {', '.join(ALLOWED_EXTENSIONS)}")

    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise ValueError(f"File exceeds max upload size of {MAX_UPLOAD_SIZE // (1024*1024)} MB")

    fd, path = tempfile.mkstemp(suffix=ext)
    with os.fdopen(fd, "wb") as f:
        f.write(content)
    return path


def parse_file(file_path: str) -> list[dict[str, str]]:
    ext = Path(file_path).suffix.lower()
    if ext == ".csv":
        return _parse_csv(file_path)
    elif ext == ".xlsx":
        return _parse_xlsx(file_path)
    raise ValueError(f"Unsupported file type: {ext}")


def _parse_csv(file_path: str) -> list[dict[str, str]]:
    with open(file_path, encoding="utf-8-sig") as f:
        sample = f.read(8192)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(f, dialect=dialect)
        rows = []
        for row in reader:
            cleaned = {k.strip(): (v.strip() if v else "") for k, v in row.items()}
            rows.append(cleaned)
        return rows


def _parse_xlsx(file_path: str) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return []

    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows_iter, [])]

    rows = []
    for row in rows_iter:
        row_dict = {}
        for idx, val in enumerate(row):
            if idx < len(headers):
                row_dict[headers[idx]] = str(val).strip() if val is not None else ""
        if any(v for v in row_dict.values()):
            rows.append(row_dict)
    wb.close()
    return rows


# ---------------------------------------------------------------------------
# Column mapping
# ---------------------------------------------------------------------------

def normalize_header(h: str) -> str:
    return re.sub(r"[_\-\s]+", " ", h.strip().lower())


def map_columns(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}

    flat_aliases: dict[str, str] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            flat_aliases[normalize_header(alias)] = field

    for header in headers:
        norm = normalize_header(header)
        if norm in flat_aliases:
            mapping[header] = flat_aliases[norm]
            continue
        best_field, best_dist = None, 3
        for alias_norm, field in flat_aliases.items():
            dist = _levenshtein(norm, alias_norm)
            if dist < best_dist:
                best_dist = dist
                best_field = field
        if best_field:
            mapping[header] = best_field

    return mapping


# ---------------------------------------------------------------------------
# Type coercion
# ---------------------------------------------------------------------------

def coerce_int(value: str, field: str) -> int | None:
    if not value:
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        raise ValueError(f"'{value}' is not a valid integer for {field}")


def coerce_list(value: str) -> list[str]:
    if not value:
        return []
    return [x.strip() for x in re.split(r"[,;|]+", value) if x.strip()]


def coerce_enum(value: str, enum_cls: type) -> Any | None:
    if not value:
        return None
    norm = value.strip().upper().replace(" ", "_")
    for member in enum_cls:
        if member.value == norm:
            return member
        if member.name == norm:
            return member
    for member in enum_cls:
        if _levenshtein(norm, member.value) < 4 or _levenshtein(norm, member.name) < 4:
            return member
    raise ValueError(f"'{value}' is not a valid {enum_cls.__name__}")


# ---------------------------------------------------------------------------
# Department resolution
# ---------------------------------------------------------------------------

async def resolve_department(
    db: AsyncSession,
    value: str,
    institution_id: uuid.UUID,
    cache: dict[str, uuid.UUID | None] | None = None,
) -> uuid.UUID | None:
    if not value:
        return None

    val_stripped = value.strip().lower()

    if cache is not None and val_stripped in cache:
        return cache[val_stripped]

    # Exact match by code
    result = await db.execute(
        select(Department).where(
            Department.code.ilike(val_stripped),
            Department.institution_id == institution_id,
            Department.is_active == True,
        )
    )
    dept = result.scalar_one_or_none()
    if dept:
        if cache is not None:
            cache[val_stripped] = dept.id
        return dept.id

    # Match by name
    result = await db.execute(
        select(Department).where(
            Department.name.ilike(val_stripped),
            Department.institution_id == institution_id,
            Department.is_active == True,
        )
    )
    dept = result.scalar_one_or_none()
    if dept:
        if cache is not None:
            cache[val_stripped] = dept.id
        return dept.id

    if cache is not None:
        cache[val_stripped] = None
    return None


# ---------------------------------------------------------------------------
# File cleanup
# ---------------------------------------------------------------------------

def cleanup_temp(file_path: str) -> None:
    try:
        if os.path.exists(file_path):
            os.unlink(file_path)
    except OSError:
        pass
